from pathlib import Path
import docx
from src.loaders.base import BaseLoader
from src.models import Document


class DocxLoader(BaseLoader):
    """
    Loader for Microsoft Word (.docx) documents.
    Extracts text paragraphs as well as tables (formatted as Markdown tables).
    """

    def __init__(self, file_path: str | Path, include_tables: bool = True):
        super().__init__(file_path)
        self.include_tables = include_tables

    def load(self) -> list[Document]:
        doc_obj = docx.Document(str(self.file_path))
        text_parts = []

        # 1. Extract paragraphs
        for para in doc_obj.paragraphs:
            text = para.text.strip()
            if text:
                text_parts.append(text)

        # 2. Extract tables if enabled
        table_count = len(doc_obj.tables)
        if self.include_tables and table_count > 0:
            for idx, table in enumerate(doc_obj.tables):
                table_lines = [f"\n--- [Table {idx + 1}] ---"]
                for row_idx, row in enumerate(table.rows):
                    row_data = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                    line = "| " + " | ".join(row_data) + " |"
                    table_lines.append(line)
                    # Add markdown header separator after first row
                    if row_idx == 0:
                        separator = "| " + " | ".join(["---"] * len(row_data)) + " |"
                        table_lines.append(separator)

                text_parts.append("\n".join(table_lines))

        full_content = "\n\n".join(text_parts)
        doc_id = f"DOC_{self.file_path.stem}"

        doc = Document(
            content=full_content,
            doc_id=doc_id,
            title=self.file_path.stem,
            source=str(self.file_path.resolve()),
            metadata={
                "file_name": self.file_path.name,
                "extension": ".docx",
                "num_paragraphs": len(doc_obj.paragraphs),
                "num_tables": table_count,
            },
        )
        return [doc]
