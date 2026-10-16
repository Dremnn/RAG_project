from pathlib import Path
from typing import Literal

from openpyxl import load_workbook

from src.loaders.base import BaseLoader
from src.models import Document


class ExcelLoader(BaseLoader):
    """
    Loader for Excel (.xlsx, .xlsm) documents using openpyxl.
    Row 1 of each sheet is treated as the header.

    split_by:
        "sheet": one Document per sheet (default)
        "row":   one Document per data row (good for FAQ / table-style data)
    """

    def __init__(
        self,
        file_path: str | Path,
        split_by: Literal["sheet", "row"] = "sheet",
    ):
        super().__init__(file_path)
        self.split_by = split_by

    @staticmethod
    def _cell_to_str(value) -> str:
        return "" if value is None else str(value).strip()

    def _row_to_text(self, headers: list[str], row: tuple) -> str:
        parts = []
        for header, value in zip(headers, row):
            text = self._cell_to_str(value)
            if text:
                parts.append(f"{header}: {text}")
        return "; ".join(parts)

    def load(self) -> list[Document]:
        wb = load_workbook(str(self.file_path), read_only=True, data_only=True)
        docs: list[Document] = []
        total_sheets = len(wb.sheetnames)

        for sheet_idx, sheet_name in enumerate(wb.sheetnames):
            ws = wb[sheet_name]
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                continue

            # Header: fall back to col_N when a header cell is empty
            headers = [
                self._cell_to_str(h) or f"col_{i + 1}"
                for i, h in enumerate(rows[0])
            ]
            data_rows = [r for r in rows[1:] if any(c is not None for c in r)]

            base_meta = {
                "file_name": self.file_path.name,
                "extension": self.file_path.suffix.lower(),
                "sheet_name": sheet_name,
                "total_sheets": total_sheets,
            }

            if self.split_by == "row":
                for row_idx, row in enumerate(data_rows, start=2):
                    text = self._row_to_text(headers, row)
                    if not text:
                        continue
                    docs.append(
                        Document(
                            content=text,
                            doc_id=f"DOC_{self.file_path.stem}_{sheet_name}_r{row_idx}",
                            title=f"{self.file_path.stem} ({sheet_name}, row {row_idx})",
                            source=str(self.file_path.resolve()),
                            metadata={**base_meta, "row_number": row_idx},
                        )
                    )
            else:
                lines = [self._row_to_text(headers, r) for r in data_rows]
                content = "\n".join(line for line in lines if line)
                if not content:
                    continue
                docs.append(
                    Document(
                        content=content,
                        doc_id=f"DOC_{self.file_path.stem}_{sheet_name}",
                        title=f"{self.file_path.stem} ({sheet_name})",
                        source=str(self.file_path.resolve()),
                        metadata={**base_meta, "total_rows": len(data_rows)},
                    )
                )

        wb.close()
        return docs